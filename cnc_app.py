"""CNC Part Analyzer - SPEC'26 (everything in one file)

  py cnc_app.py                      -> start the app (asks for part number + questions)
  py cnc_app.py --check 2280         -> show files, facts and cross-check for one part (no AI)
  py cnc_app.py --data D:\\folder    -> use files from another folder (e.g. the judges' USB)
  py cnc_app.py --readme             -> write README.md

Put the Excel (.xlsx/.csv), PDF drawings (.pdf) and IGES models (.igs/.iges) in the 'data' folder.
"""
import csv
import json
import math
import re
import sys
import urllib.request
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook, Workbook
from pypdf import PdfReader

HERE = Path(__file__).parent
DATA = HERE / "data"
KEY_FILE = HERE / "llm_key.txt"
LLM_URL = "https://api.groq.com/openai/v1"
if "--data" in sys.argv[:-1]:
    DATA = Path(sys.argv[sys.argv.index("--data") + 1])

try:
    sys.stdout.reconfigure(errors="replace")
except AttributeError:
    pass


def clean(value):
    if value is None:
        return ""
    text = str(value).strip()
    return text[:-2] if text.endswith(".0") else text


# =====================================================================
# 1. EXCEL: find the part's row in the manufacturing record
# =====================================================================
def sheet_rows(path):
    """Yield (sheet name, list of rows) for .xlsx and .csv files."""
    if path.suffix.lower() == ".csv":
        with open(path, newline="", encoding="utf-8-sig", errors="ignore") as f:
            yield "csv", [tuple(r) for r in csv.reader(f)]
        return
    wb = load_workbook(path, read_only=True, data_only=True)
    for ws in wb.worksheets:
        yield ws.title, list(ws.iter_rows(values_only=True))
    wb.close()


def find_header(rows):
    """The column-title row: the first of the top 15 rows with 3+ text cells (titles can be below row 1)."""
    for i, row in enumerate(rows[:15]):
        texts = [v for v in row if clean(v) and num(clean(v)) is None]
        if len(texts) >= 3 and len(texts) >= len([v for v in row if clean(v)]) * 0.6:
            return i
    return 0


def part_in(value, part, exact):
    v = clean(value)
    if exact:
        return v.lower() == part.lower()
    return re.search(rf"(?<![\w.]){re.escape(part)}(?![\w.])", v, re.I) is not None


def read_excel(path, part):
    """Find every row that contains the part number, on every sheet.
    First tries an exact cell match; if nothing, accepts cells like 'P-2280' or '2280-A'."""
    sheets = list(sheet_rows(path))
    for exact in (True, False):
        facts = []
        for name, rows in sheets:
            if not rows:
                continue
            h = find_header(rows)
            headers = [clean(x) or f"Column {i + 1}" for i, x in enumerate(rows[h])]
            for n, row in enumerate(rows[h + 1:], start=h + 2):
                if any(part_in(v, part, exact) for v in row):
                    for col, val in zip(headers, row):
                        if clean(val):
                            facts.append({"source": f"Excel: {path.name}, sheet '{name}', row {n}",
                                          "field": col, "value": clean(val)})
        if facts:
            return facts
    return []


# =====================================================================
# 2. PDF: read the text and numbers on the engineering drawing
# =====================================================================
KEYWORDS = ["MATERIAL", "MATL", "FINISH", "TOLERANCE", "SCALE", "UNIT", "MM", "INCH", "PART",
            "DWG", "DRAWING", "REV", "TITLE", "QTY", "THREAD", "HOLE", "DEPTH", "SURFACE", "GRADE",
            "ALLOY", "HEAT", "ANODIZE", "COAT", "DRAWN", "DATE", "SHEET", "WEIGHT", "NOTE",
            "CHAMFER", "FILLET", "RADIUS", "DIA", "DEBURR", "ROUGHNESS"]


def pdf_text(path):
    return [page.extract_text() or "" for page in PdfReader(path).pages]


def read_pdf(path):
    facts, numbers = [], []
    for page_no, text in enumerate(pdf_text(path), start=1):
        if not text.strip():
            facts.append({"source": f"PDF drawing: {path.name}, page {page_no}", "field": "warning",
                          "value": "No readable text on this page (scanned image?)"})
            continue
        lines, pending = [], ""
        for line in text.splitlines():            # join "MATERIAL:" with the value on the next line
            line = line.strip()
            if not line:
                continue
            if line.endswith(":") and len(line) < 30:
                pending = (pending + " " + line).strip()
                continue
            lines.append((pending + " " + line).strip())
            pending = ""
        if pending:
            lines.append(pending)
        for line in lines:
            if line:
                important = any(k in line.upper() for k in KEYWORDS)
                facts.append({"source": f"PDF drawing: {path.name}, page {page_no}",
                              "field": "title block / note" if important else "drawing text",
                              "value": line})
                numbers += [float(x) for x in re.findall(r"\d+(?:\.\d+)?", line)]
    notes = path.with_name(path.stem + "_notes.txt")   # backup for scanned drawings
    if notes.exists():
        for line in notes.read_text(encoding="utf-8").splitlines():
            if line.strip():
                facts.append({"source": f"PDF drawing: {path.name} (typed from drawing by team)",
                              "field": "title block / note", "value": line.strip()})
                numbers += [float(x) for x in re.findall(r"\d+(?:\.\d+)?", line)]
    return facts, numbers


# =====================================================================
# 3. IGES: units, overall size and circles from the 3D model
# =====================================================================
UNITS = {1: "inch", 2: "mm", 4: "ft", 6: "m", 7: "km", 8: "mil", 9: "micron", 10: "cm"}
SHAPES = {100: "circular arc", 102: "composite curve", 110: "line", 116: "point", 120: "surface of revolution",
          122: "tabulated cylinder", 124: "transformation", 126: "B-spline curve", 128: "B-spline surface",
          141: "boundary", 142: "curve on surface", 143: "bounded surface", 144: "trimmed surface",
          108: "plane", 123: "direction", 186: "solid", 190: "plane", 192: "cylinder", 314: "color", 406: "property",
          502: "vertex list", 504: "edge list", 508: "loop", 510: "face", 514: "shell"}


def split_params(s, pd=",", rd=";"):
    fields, i, n = [], 0, len(s)
    while i < n:
        m = re.match(r"\s*(\d+)H", s[i:])
        if m:
            k = int(m.group(1))
            start = i + m.end()
            fields.append(s[start:start + k])
            i = start + k
        else:
            j = i
            while j < n and s[j] not in (pd, rd):
                j += 1
            fields.append(s[i:j].strip())
            i = j
        if i < n and s[i] == rd:
            break
        i += 1
    return fields


def num(x):
    try:
        return float(str(x).strip().replace("D", "E"))
    except ValueError:
        return None


def bspline_points(p, samples=24):
    """Evaluate a (rational) B-spline curve (IGES 126) at evenly spaced parameters."""
    K, M = int(p[1]), int(p[2])
    n = K + 1
    knots = p[7:7 + n + M + 1]
    w = p[7 + n + M + 1:7 + n + M + 1 + n]
    c0 = 7 + n + M + 1 + n
    cps = [p[c0 + 3 * j:c0 + 3 * j + 3] for j in range(n)]
    v0, v1 = p[c0 + 3 * n], p[c0 + 3 * n + 1]
    out = []
    for k in range(samples + 1):
        t = v0 + (v1 - v0) * k / samples
        span = M
        while span < n - 1 and knots[span + 1] <= t:
            span += 1
        d = [[cps[j][0] * w[j], cps[j][1] * w[j], cps[j][2] * w[j], w[j]] for j in range(span - M, span + 1)]
        for r in range(1, M + 1):
            for j in range(M, r - 1, -1):
                i = span - M + j
                den = knots[i + M + 1 - r] - knots[i]
                a = 0.0 if den == 0 else (t - knots[i]) / den
                d[j] = [(1 - a) * d[j - 1][q] + a * d[j][q] for q in range(4)]
        x, y, z, ww = d[M]
        out.append((x / ww, y / ww, z / ww) if ww else (x, y, z))
    return out


def read_iges(path):
    """Read units, overall size and circles from an IGES file.
    Real CAD files place shapes with transformation matrices (entity 124), so we apply them."""
    src = f"IGES 3D model: {path.name}"
    sec = {"S": [], "G": [], "D": [], "P": []}
    for line in path.read_text(errors="ignore").splitlines():
        line = line.ljust(80)
        if line[72] in sec:
            sec[line[72]].append(line)
    g = "".join(l[:72] for l in sec["G"])
    pd, rd = (g[2] if g.startswith("1H") else ","), ";"
    gf = split_params(g, pd, rd)
    if len(gf) > 1 and len(gf[1]) == 1:
        rd = gf[1]
    unit = "unknown"
    if len(gf) > 14:
        flag = num(gf[13])
        unit = UNITS.get(int(flag), gf[14] or "unknown") if flag else (gf[14] or "unknown")
    ptext = {}
    for l in sec["P"]:
        try:
            ptext[int(l[64:72])] = ptext.get(int(l[64:72]), "") + l[:64]
        except ValueError:
            pass

    # Directory entries: number -> (type, transform pointer, status)
    ents, d = {}, sec["D"]
    for i in range(0, len(d) - 1, 2):
        try:
            ents[i + 1] = (int(d[i][0:8]), int(d[i][48:56] or 0), d[i][64:72].replace(" ", "0"))
        except ValueError:
            continue
    params = {de: [num(x) for x in split_params(ptext.get(de, ""), pd, rd)] for de in ents}

    def transform(pt, tp, depth=0):
        """Apply transformation matrix entity 124 (and any parent matrices)."""
        if not tp or tp not in ents or ents[tp][0] != 124 or depth > 10:
            return pt
        m = params[tp][1:13]
        x, y, z = pt
        new = (m[0] * x + m[1] * y + m[2] * z + m[3],
               m[4] * x + m[5] * y + m[6] * z + m[7],
               m[8] * x + m[9] * y + m[10] * z + m[11])
        return transform(new, ents[tp][1], depth + 1)

    # Real CAD files also contain helper curves (e.g. the axis of a round face). The real edges are
    # listed by edge lists (504, solid files) or curve-on-surface entities (142, surface files).
    edge_curves = set()

    def add_curve(de, depth=0):
        if de not in ents or depth > 5:
            return
        edge_curves.add(de)
        if ents[de][0] == 102:                       # composite curve -> its pieces
            q = params[de]
            for k in range(int(q[1])):
                add_curve(int(q[2 + k]), depth + 1)

    for de, (etype, _, _) in ents.items():
        q = params[de]
        try:
            if etype == 504:
                for k in range(int(q[1])):
                    add_curve(int(q[2 + 5 * k]))
            elif etype == 142 and q[4]:
                add_curve(int(q[4]))
        except (TypeError, ValueError, IndexError):
            pass

    counts, curve_pts, other_pts, circles = Counter(), [], [], []
    for de, (etype, tp, status) in ents.items():
        counts[etype] += 1
        if status[4:6] == "05":   # 2D parameter-space curve, not real geometry
            continue
        if edge_curves and etype in (110, 100, 126) and de not in edge_curves:
            continue
        p = params[de]
        try:
            if etype == 110:                                   # line
                curve_pts += [transform(tuple(p[1:4]), tp), transform(tuple(p[4:7]), tp)]
            elif etype == 100:                                 # circular arc / circle
                zt, cx, cy, sx, sy, ex, ey = p[1:8]
                r = math.hypot(sx - cx, sy - cy)
                a0, a1 = math.atan2(sy - cy, sx - cx), math.atan2(ey - cy, ex - cx)
                full = abs(sx - ex) < 1e-6 and abs(sy - ey) < 1e-6
                sweep = 2 * math.pi if full else (a1 - a0) % (2 * math.pi)
                curve_pts += [transform((cx + r * math.cos(a0 + sweep * k / 36),
                                         cy + r * math.sin(a0 + sweep * k / 36), zt), tp) for k in range(37)]
                if full:
                    centre = tuple(round(v, 2) for v in transform((cx, cy, zt), tp))
                    circles.append((round(2 * r, 3), centre))
            elif etype == 126:                                 # B-spline curve
                curve_pts += [transform(q, tp) for q in bspline_points(p)]
            elif etype == 502:                                 # vertex list (solid models)
                other_pts += [transform(tuple(p[2 + 3 * j:5 + 3 * j]), tp) for j in range(int(p[1]))]
            elif etype == 116 and status[2:4] == "00":         # stand-alone point
                other_pts.append(transform(tuple(p[1:4]), tp))
            elif etype == 128:                                 # B-spline surface (fallback only)
                K1, K2, M1, M2 = int(p[1]), int(p[2]), int(p[3]), int(p[4])
                n = (K1 + 1) * (K2 + 1)
                s = 10 + (K1 + M1 + 2) + (K2 + M2 + 2) + n
                other_pts += [transform(tuple(p[s + 3 * j:s + 3 * j + 3]), tp) for j in range(n)]
        except (TypeError, ValueError, IndexError, ZeroDivisionError):
            continue
    pts = [q for q in (curve_pts or other_pts) if len(q) == 3 and None not in q]
    holes = sorted({dia for dia, _ in circles})
    facts = [{"source": src, "field": "units", "value": unit}]
    geo = {"units": unit, "size": None, "holes": holes}
    if pts:
        xs, ys, zs = zip(*pts)
        geo["size"] = tuple(round(max(a) - min(a), 2) for a in (xs, ys, zs))
        facts.append({"source": src, "field": "overall size (bounding box X x Y x Z)",
                      "value": " x ".join(f"{v:g}" for v in geo["size"]) + f" {unit}"})
    else:
        facts.append({"source": src, "field": "overall size",
                      "value": "could not be measured (this file uses shape types our reader does not support)"})
    if circles:
        for dia in holes:
            n = len({c for d_, c in circles if d_ == dia})
            facts.append({"source": src, "field": "full circle edges (often holes or round faces)",
                          "value": f"diameter {dia:g} {unit}: {n} circle edge(s) "
                                   f"(a through-hole usually has 2, one at each end)"})
    facts.append({"source": src, "field": "shapes in file",
                  "value": ", ".join(f"{c} x {SHAPES.get(t, 'type ' + str(t))}" for t, c in counts.most_common(8))})
    return facts, geo


# =====================================================================
# 4. CROSS-CHECK: compare the three files (normal code, not the AI)
# =====================================================================
MATERIAL_WORDS = ["material", "matl", "mat'l", "grade", "alloy"]
SIZE_WORDS = ["length", "width", "height", "thickness", "diameter", "dia", "depth", "od", "id"]
TO_MM = {"mm": 1.0, "inch": 25.4, "cm": 10.0, "m": 1000.0}


def words_of(text):
    return re.findall(r"[a-z0-9]+", text.lower())


def cross_check(part, excel, pdf, pdf_numbers, geo, tol=0.5):
    """Compare the three files. Returns a list of (status, message). Done in Python, not by the AI."""
    res = []
    if pdf:
        ok = any(re.search(rf"(?<!\d){re.escape(part)}(?!\d)", f["value"]) for f in pdf)
        res.append(("OK" if ok else "CHECK", f"Part number {part} {'appears' if ok else 'was NOT found'} on the PDF drawing."))

    # Material: Excel vs drawing
    mat = next((f for f in excel if any(w in f["field"].lower() for w in MATERIAL_WORDS)), None)
    if mat and pdf:
        dwg_mat = ""
        for i, f in enumerate(pdf):
            if "MATERIAL" in f["value"].upper() or "MATL" in f["value"].upper():
                dwg_mat = f["value"]
                if dwg_mat.rstrip().endswith(":") and i + 1 < len(pdf):   # label and value on separate lines
                    dwg_mat += " " + pdf[i + 1]["value"]
                break
        want = [w for w in words_of(mat["value"]) if w not in ("t6", "t4", "aisi")][:2]   # e.g. aluminium 6061
        pdf_words = set(words_of(" ".join(f["value"] for f in pdf)))
        if want and all(w in pdf_words for w in want):
            res.append(("OK", f"Material '{mat['value']}' is the same in Excel and the PDF drawing."))
        else:
            shown = f" The drawing says: '{dwg_mat}'." if dwg_mat else ""
            res.append(("MISMATCH", f"Excel says material is '{mat['value']}' but the drawing does not.{shown}"))

    # Sizes: Excel vs drawing
    excel_sizes = []
    for f in excel:
        words = f["field"].lower().replace("(", " ").replace(")", " ").replace("_", " ").split()
        v = num(f["value"])
        if any(w in SIZE_WORDS for w in words) and v is not None:
            excel_sizes.append(v)
            if pdf_numbers:
                ok = any(abs(n - v) <= tol for n in pdf_numbers)
                res.append(("OK" if ok else "MISMATCH",
                            f"Excel {f['field']} = {v:g} {'matches a dimension on' if ok else 'has NO match on'} the drawing."))

    # Sizes: 3D model vs drawing and vs Excel
    if geo and geo["size"]:
        factor = TO_MM.get(geo["units"], 1.0)
        size_mm = [v * factor for v in geo["size"]]
        if pdf_numbers:
            for axis, v in zip("XYZ", size_mm):
                if v > 0:
                    ok = any(abs(n - v) <= tol for n in pdf_numbers)
                    res.append(("OK" if ok else "CHECK",
                                f"3D model {axis}-size {v:g} mm {'matches a dimension on' if ok else 'has no match on'} the drawing."))
        if len(excel_sizes) >= 2:
            a_, b_ = sorted(excel_sizes, reverse=True), sorted(size_mm, reverse=True)
            ok = all(abs(x - y) <= tol for x, y in zip(a_, b_))
            res.append(("OK" if ok else "MISMATCH",
                        f"Excel sizes {' x '.join(f'{x:g}' for x in a_)} {'match' if ok else 'do NOT match'} "
                        f"the 3D model size {' x '.join(f'{y:g}' for y in b_)} mm."))
    if geo and pdf_numbers:
        for dia in geo["holes"]:
            ok = any(abs(n - dia) <= 0.05 for n in pdf_numbers)
            res.append(("OK" if ok else "CHECK",
                        f"Circle diameter {dia:g} in the 3D model {'is' if ok else 'is NOT'} on the drawing."))
    return res


# =====================================================================
# 5. AI (LLM): answer only from the facts, with sources
# =====================================================================
RULES = """You are a CNC manufacturing assistant.
You get FACTS about one part. Each fact starts with its source in [square brackets].
Rules:
1. Answer ONLY from the FACTS. Never use outside knowledge. Never guess numbers.
2. After every claim, write its source in [square brackets], exactly as given.
3. If the FACTS do not contain the answer, say: "I can't find that in the files." and say briefly what IS available.
4. If the CROSS-CHECK shows a MISMATCH or CHECK related to the question, mention it.
5. Keep it short and simple. Always give units."""


def api(path, body=None, key=""):
    req = urllib.request.Request(LLM_URL + path, data=json.dumps(body).encode() if body else None,
                                 headers={"Authorization": f"Bearer {key}", "Content-Type": "application/json",
                                          "User-Agent": "Mozilla/5.0"})
    with urllib.request.urlopen(req, timeout=60) as r:
        return json.loads(r.read().decode())


def get_key():
    if KEY_FILE.exists() and KEY_FILE.read_text().strip():
        return KEY_FILE.read_text().strip()
    print("\nYou need a free Groq API key (get one at console.groq.com -> API Keys).")
    key = input("Paste your key here and press Enter: ").strip()
    KEY_FILE.write_text(key)
    return key


def pick_model(key):
    ids = [m["id"] for m in api("/models", key=key)["data"]]
    chat = [i for i in ids if not any(b in i.lower() for b in ("whisper", "guard", "tts", "embed", "audio"))]
    for want in ("llama-3.3-70b", "70b", "llama"):
        for i in chat:
            if want in i:
                return i
    return chat[0]


def ask_ai(key, model, part, question, facts, checks):
    fact_lines = "\n".join(f"- [{f['source']}] {f['field']}: {f['value']}" for f in facts)
    check_lines = "\n".join(f"- {s}: {m}" for s, m in checks) or "- none"
    prompt = f"PART: {part}\n\nFACTS:\n{fact_lines}\n\nCROSS-CHECK:\n{check_lines}\n\nQUESTION: {question}"
    out = api("/chat/completions", {"model": model, "temperature": 0,
                                     "messages": [{"role": "system", "content": RULES},
                                                  {"role": "user", "content": prompt}]}, key)
    return out["choices"][0]["message"]["content"]


# =====================================================================
# 6. FIND THE FILES FOR A PART (works for any part number)
# =====================================================================
def mentions(text, part):
    return re.search(rf"(?<!\d){re.escape(part)}(?!\d)", text) is not None


def pick(kind, files, part, excel_facts):
    """Choose the file for this part. Order: 1) name has the part number, 2) the Excel row names the file,
    3) the part number is inside the file, 4) it is the only file of its kind. Returns (file, how)."""
    for f in files:
        if mentions(f.stem, part):
            return f, "file name contains the part number"
    named = {Path(x["value"]).name.lower() for x in excel_facts}
    for f in files:
        if f.name.lower() in named:
            return f, "file is named in the Excel record"
    for f in files:
        try:
            text = " ".join(pdf_text(f)) if kind == "pdf" else f.read_text(errors="ignore")[:30000]
        except Exception:
            continue
        if mentions(text, part):
            return f, "part number found inside the file"
    if len(files) == 1:
        return files[0], "only file of this type in the folder (NOT confirmed for this part)"
    return None, ""


def load_part(part, show=True):
    files = sorted(p for p in DATA.iterdir() if p.is_file()) if DATA.exists() else []
    excels = [p for p in files if p.suffix.lower() in (".xlsx", ".xlsm", ".csv") and not p.name.startswith("~$")]
    excel = []
    for x in excels:
        try:
            excel += read_excel(x, part)
        except Exception as e:
            print(f"  !! Could not read {x.name}: {e}")
    pdf_file, pdf_how = pick("pdf", [p for p in files if p.suffix.lower() == ".pdf"], part, excel)
    igs_file, igs_how = pick("igs", [p for p in files if p.suffix.lower() in (".igs", ".iges")], part, excel)
    pdf, pdf_numbers, iges, geo = [], [], [], None
    if pdf_file:
        pdf, pdf_numbers = read_pdf(pdf_file)
    if igs_file:
        iges, geo = read_iges(igs_file)
    checks = cross_check(part, excel, pdf, pdf_numbers, geo)
    notes = []
    for label, f, how in (("PDF drawing", pdf_file, pdf_how), ("IGES 3D model", igs_file, igs_how)):
        if f and "NOT confirmed" in how:
            notes.append({"source": "App", "field": "warning", "value": f"{label} {f.name} was used because it is the {how}"})
    for label, got in (("Excel row", excel), ("PDF drawing", pdf_file), ("IGES 3D model", igs_file)):
        if not got:
            notes.append({"source": "App", "field": "missing file", "value": f"No {label} was found for part {part}"})
    if show:
        print(f"\n  Files used for part {part}  (folder: {DATA})")
        rows = sorted({f["source"] for f in excel})
        print(f"    Excel: {len(rows)} row(s) found" if rows else "    Excel: part NOT found in Excel")
        for r in rows:
            print(f"           - {r}")
        print(f"    PDF:   {pdf_file.name + '  <- ' + pdf_how if pdf_file else 'not found'}")
        print(f"    IGES:  {igs_file.name + '  <- ' + igs_how if igs_file else 'not found'}")
        print("\n  Cross-check:")
        for st, m in checks or [("-", "nothing to compare")]:
            print(f"    [{st}] {m}")
    if not (excel or pdf_file or igs_file):
        return [], checks
    return excel + pdf + iges + notes, checks


STOP = set("what is the of a an this part and are for in to does it how much many which do has have with on by its".split())


def offline_answer(question, facts):
    """Backup when the AI is not reachable: show the facts that contain the question's words."""
    words = [w for w in words_of(question) if w not in STOP and len(w) > 2]
    hits = [f for f in facts if any(w in (f["field"] + " " + f["value"]).lower() for w in words)]
    if not hits:
        return "I can't find that in the files."
    return "(AI offline - matching facts)\n" + "\n".join(f"- {f['field']}: {f['value']}  [{f['source']}]" for f in hits[:12])


README = """# CNC Part Analyzer (SPEC'26)

Type a part number and a question. The app finds that part's Excel manufacturing record,
PDF engineering drawing and IGES 3D model, cross-checks them, and asks an LLM to answer
using ONLY those facts. Every answer names the file (and sheet/row/page) it came from.
If the answer is not in the files, it says "I can't find that in the files."

## How to run (Windows)
1. Install Python 3.10+ from python.org.
2. Install the two libraries:   py -m pip install openpyxl pypdf
3. Put the Excel (.xlsx or .csv), PDF drawings (.pdf) and IGES models (.igs/.iges) in the `data` folder.
4. Run:   py cnc_app.py
5. First time only: paste a free Groq API key (console.groq.com -> API Keys). It is saved in llm_key.txt.
6. Type a part number (e.g. 2280), then ask questions.
   Commands: `facts` = show all evidence, `new` = another part, `quit` = exit.

Other commands:
- py cnc_app.py --check 2280        show files found, facts and cross-check (no AI needed)
- py cnc_app.py --data D:\\files    use files from another folder (e.g. new files from the judges)

## Sample dataset (data folder)
5 parts (2280-2284): one Excel record (2 sheets), 5 PDF drawings, 5 IGES models. See data/SOURCES.txt.
Part 2283 has a deliberate material mismatch; part 2282 has no surface finish in Excel (honesty test).

## How it works
1. Excel (openpyxl): finds the column-title row automatically, then every row on every sheet
   that contains the part number (exact match first, then forms like P-2280). Keeps sheet + row.
2. Finding the PDF/IGES for a part: file name contains the part number -> file named in the
   Excel row -> part number inside the file -> only file of that type (shown as a warning).
3. PDF (pypdf): text lines and numbers from every page; title-block lines are marked.
4. IGES (own parser, no CAD software needed): units from the Global section; real edges
   (lines, arcs, B-spline curves) with transformation matrices applied give the overall size;
   full circles give hole/round diameters.
5. Cross-check (plain Python, not the AI): part number on drawing, material Excel vs drawing,
   Excel sizes vs drawing, 3D size vs drawing, Excel sizes vs 3D size, circle diameters vs drawing.
   Results: OK / MISMATCH / CHECK.
6. LLM: Groq API (OpenAI-compatible); the model is chosen automatically and printed at start.
   Temperature 0. The LLM only receives the extracted facts with their sources plus strict
   rules (answer only from facts, cite sources, say when something is missing).
   If the AI cannot be reached, an offline mode shows the matching facts instead.

## Limitations
- Size is the bounding box along the model's X/Y/Z axes (a rotated part gives a bigger box).
- Circles are reported as "circle edges", not guaranteed holes (could be round outer faces).
- Old IGES types (e.g. 112/114 parametric splines) are not measured; the app says so.
- Scanned PDFs have no text; facts can be typed into <pdf name>_notes.txt as a workaround.
- Drawing numbers are matched without knowing which dimension line they belong to.
- .xls (old Excel) is not supported: save as .xlsx or .csv.
- The LLM can still misread a fact, so the app always shows the evidence it was given.
"""


# =====================================================================
# 8. THE APP
# =====================================================================
def main():
    if "--readme" in sys.argv:
        (HERE / "README.md").write_text(README, encoding="utf-8")
        return print("Wrote README.md")
    if "--check" in sys.argv[:-1]:
        facts, _ = load_part(sys.argv[sys.argv.index("--check") + 1])
        print("\n  Facts:")
        for f in facts:
            print(f"    [{f['source']}] {f['field']}: {f['value']}")
        return
    DATA.mkdir(exist_ok=True)
    print("=" * 60 + "\n  CNC PART ANALYZER  (Excel + PDF drawing + IGES 3D model)\n" + "=" * 60)
    print(f"  Data folder: {DATA}")
    key = get_key()
    try:
        model = pick_model(key)
        print(f"  AI model: {model}")
    except Exception as e:
        print(f"  !! Could not reach the AI ({e}).\n     Check internet / the key in llm_key.txt. Using offline mode for now.")
        model = None
    while True:
        part = input("\nPart number (or 'quit'): ").strip()
        if part.lower() in ("quit", "q", "exit"):
            break
        if not part:
            continue
        facts, checks = load_part(part)
        if not facts:
            print(f"\n  I can't find part {part} in any of the files.")
            continue
        while True:
            q = input(f"\nQuestion about {part} ('new' = other part, 'facts' = show evidence, 'quit'): ").strip()
            if q.lower() in ("quit", "q", "exit"):
                return
            if q.lower() == "new":
                break
            if q.lower() == "facts":
                for f in facts:
                    print(f"  [{f['source']}] {f['field']}: {f['value']}")
                continue
            if not q:
                continue
            if model:
                try:
                    print("\nAnswer:\n" + ask_ai(key, model, part, q, facts, checks))
                    continue
                except Exception as e:
                    print(f"  !! AI error: {e}")
            print("\nAnswer:\n" + offline_answer(q, facts))


if __name__ == "__main__":
    main()
